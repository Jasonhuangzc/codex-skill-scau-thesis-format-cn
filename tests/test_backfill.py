"""Cross-platform regressions for content-preserving chapter/pipeline backfill.

Run with: python -m unittest discover -s tests -p 'test_backfill.py' -v
Only python-docx (and its lxml dependency) are required; no Word COM/assets.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SCRIPTS = Path(__file__).resolve().parents[1] / 'scau-thesis-format-cn' / 'scripts'
sys.path.insert(0, str(SCRIPTS))

from insert_markdown_chapter import parse_markdown, required_donor_keys
from run_scau_project_pipeline import strip_template_body_placeholders
from word_template_utils import (
    delete_range, find_heading_donors, insert_paragraph_after,
    iter_block_items, set_east_asia_font,
)


class BackfillRegression(unittest.TestCase):
    def setUp(self):
        self.workspace = tempfile.TemporaryDirectory(prefix='scau-backfill-tests-')
        self.root = Path(self.workspace.name)
        self.source = self.root / 'working.docx'
        self.markdown = self.root / 'chapter.md'
        self.output = self.root / 'result.docx'

    def tearDown(self):
        self.workspace.cleanup()

    def make_document(self, *, all_donors=False):
        document = Document()
        document.add_paragraph('封面不应成为正文 donor')
        first = document.add_paragraph('1  绪论', style='Heading 1')
        set_east_asia_font(first.runs[0], east_asia_font='黑体', ascii_font='Times New Roman', size_pt=14)
        document.add_paragraph('真实第一章正文，应一直保留。')
        document.add_paragraph('1.1  背景', style='Heading 2')
        if all_donors:
            document.add_paragraph('1.1.1  三级标题', style='Heading 3')
            document.add_paragraph('1.1.1.1  四级标题', style='Heading 4')
            document.add_paragraph('表1-1  样例表题')
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = '模板表格'
            document.add_paragraph('注：样例注释')
        document.add_paragraph('2  原章节标题', style='Heading 1')
        document.add_paragraph('旧第二章正文')
        document.add_paragraph('2.1  原小节', style='Heading 2')
        document.add_paragraph('旧第二章小节内容')
        document.add_paragraph('参  考  文  献', style='Heading 1')
        document.add_paragraph('参考文献原条目')
        document.add_paragraph('附录A  材料', style='Heading 1')
        document.add_paragraph('附录内容')
        document.add_paragraph('致        谢', style='Heading 1')
        document.add_paragraph('致谢原内容')
        document.save(self.source)
        return document

    def write_chapter(self, text='# 第2章 新章节标题\n\n## 2.1 新小节\n\n修订后的第二章内容。\n'):
        self.markdown.write_text(text, encoding='utf-8')

    def chapter_command(self, source=None, output=None):
        return [sys.executable, str(SCRIPTS / 'insert_markdown_chapter.py'), '--docx', str(source or self.source), '--chapter-file', str(self.markdown), '--output', str(output or self.output)]

    def pipeline_command(self, **kwargs):
        command = [sys.executable, str(SCRIPTS / 'run_scau_project_pipeline.py'), '--project-root', str(self.root), '--docx', str(kwargs.get('source', self.source)), '--output', str(kwargs.get('output', self.output)), '--skip-frontmatter', '--skip-figures', '--skip-tables']
        if kwargs.get('skip_chapter'):
            command.append('--skip-chapter')
        else:
            command.extend(['--chapter-file', str(self.markdown)])
        if kwargs.get('references'):
            command.extend(['--references-file', str(kwargs['references'])])
        else:
            command.append('--skip-references')
        return command

    def run_cli(self, command, expected=0):
        result = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, expected, result.stderr + result.stdout)
        return result

    def text_blocks(self, path):
        return [block.text if isinstance(block, Paragraph) else [[cell.text for cell in row.cells] for row in block.rows] for block in iter_block_items(Document(path))]

    def test_chapter_only_does_not_need_metadata_or_template_and_preserves_other_sections(self):
        self.make_document()
        self.write_chapter()
        original = self.source.read_bytes()
        result = self.run_cli(self.pipeline_command())
        report = json.loads(result.stdout)
        text = self.text_blocks(self.output)
        self.assertIn('真实第一章正文，应一直保留。', text)
        self.assertIn('附录内容', text)
        self.assertIn('参考文献原条目', text)
        self.assertIn('致谢原内容', text)
        self.assertNotIn('旧第二章正文', text)
        self.assertNotIn('2  原章节标题', text)
        self.assertEqual(text.count('2  新章节标题'), 1)
        self.assertEqual(report['steps_run'], ['chapter'])
        self.assertFalse(report['template_body_trimmed'])
        self.assertTrue(report['contents_refresh_required'])
        self.assertEqual(self.source.read_bytes(), original)

    def test_second_backfill_is_idempotent_even_when_title_changed(self):
        self.make_document()
        self.write_chapter()
        self.run_cli(self.chapter_command())
        second = self.root / 'second.docx'
        self.run_cli(self.chapter_command(self.output, second))
        self.assertEqual(self.text_blocks(self.output), self.text_blocks(second))

    def add_plain_toc_field(self, document):
        anchor = document.paragraphs[0]
        elements = []
        for text in ('', '2  原章节标题', '参  考  文  献'):
            element = OxmlElement('w:p')
            anchor._p.addprevious(element)
            paragraph = Paragraph(element, document)
            if text:
                paragraph.add_run(text)
            elements.append(paragraph)
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        elements[0].add_run()._r.append(begin)
        instruction = OxmlElement('w:instrText')
        instruction.text = ' TOC \\o "1-3" '
        elements[0].add_run()._r.append(instruction)
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        elements[-1].add_run()._r.append(end)
        document.save(self.source)

    def test_multiline_unstyled_toc_does_not_supply_chapter_or_reference_anchor(self):
        document = self.make_document()
        self.add_plain_toc_field(document)
        self.write_chapter()
        self.run_cli(self.chapter_command())
        text = self.text_blocks(self.output)
        self.assertEqual(text.count('2  原章节标题'), 1)  # preserved TOC result only
        self.assertEqual(text.count('2  新章节标题'), 1)
        self.assertIn('真实第一章正文，应一直保留。', text)

        self.write_chapter('# 第3章 新增章\n\n新第三章内容。\n')
        self.run_cli(self.chapter_command())
        text = self.text_blocks(self.output)
        self.assertGreater(text.index('3  新增章'), text.index('2  原章节标题', 4))
        self.assertEqual(text[text.index('新第三章内容。') + 1], '参  考  文  献')

    def test_number_at_start_of_body_paragraph_is_not_a_next_chapter_boundary(self):
        document = self.make_document()
        old_body = document.paragraphs[5]
        old_body.text = '4 次处理后的旧实验说明'
        document.save(self.source)
        self.write_chapter()
        self.run_cli(self.chapter_command())
        self.assertNotIn('4 次处理后的旧实验说明', self.text_blocks(self.output))
        self.assertNotIn('旧第二章小节内容', self.text_blocks(self.output))

    def test_tab_between_heading_number_and_title_is_not_a_contents_entry(self):
        document = self.make_document()
        document.paragraphs[4].text = '2\t原章节标题'
        document.save(self.source)
        self.write_chapter()
        self.run_cli(self.chapter_command())
        text = self.text_blocks(self.output)
        self.assertNotIn('2\t原章节标题', text)
        self.assertNotIn('旧第二章小节内容', text)
        self.assertEqual(text.count('2  新章节标题'), 1)

    def test_ambiguous_direct_format_numbered_boundaries_fail_closed(self):
        document = self.make_document()
        document.paragraphs[4].style = 'Normal'
        document.paragraphs[5].text = '4 次处理组说明'
        document.save(self.source)
        self.write_chapter()
        result = self.run_cli(self.chapter_command(), expected=1)
        self.assertIn('Ambiguous numbered paragraph', result.stderr)
        self.assertFalse(self.output.exists())

    def test_media_and_missing_table_payload_require_explicit_rebuild_permission(self):
        self.write_chapter()
        for case, flag in [('media', '--replace-media'), ('table', '--replace-tables')]:
            with self.subTest(case=case):
                self.output.unlink(missing_ok=True)
                document = self.make_document()
                if case == 'media':
                    document.paragraphs[5].add_run()._r.append(OxmlElement('w:drawing'))
                else:
                    table = document.add_table(rows=1, cols=1)
                    table.cell(0, 0).text = '原始数据不可默默删除'
                    document.paragraphs[5]._p.addnext(table._tbl)
                document.save(self.source)
                original = self.source.read_bytes()
                failed = self.run_cli(self.pipeline_command(), expected=1)
                recovery = json.loads(failed.stderr[failed.stderr.index('{'):])
                shutil.rmtree(recovery['recovery_dir'])
                self.assertFalse(self.output.exists())
                self.assertEqual(self.source.read_bytes(), original)
                self.run_cli(self.pipeline_command() + [flag])
                self.assertEqual(self.source.read_bytes(), original)
                self.assertFalse(Document(self.output).element.xpath('.//w:drawing | .//w:tbl'))

    def test_footnotes_equations_and_objects_have_no_structural_replacement_bypass(self):
        self.write_chapter()
        for tag in ('w:footnoteReference', 'm:oMath', 'w:object'):
            with self.subTest(tag=tag):
                document = self.make_document()
                document.paragraphs[5].add_run()._r.append(OxmlElement(tag))
                document.save(self.source)
                self.run_cli(self.chapter_command() + ['--replace-media', '--replace-tables'], expected=1)
                self.assertFalse(self.output.exists())

    def test_table_caption_continuation_note_and_rows_are_in_school_order(self):
        self.make_document(all_donors=True)
        rows = '\n'.join(f'{i} | value {i}' for i in range(1, 28))
        self.write_chapter('# 第2章 表格回灌\n\n## 2.1 结果\n\n正文。\n\n表2-1  长表\n\n编号 | 名称\n--- | ---\n' + rows + '\n\n注：表下注释\n\n表后正文。\n')
        self.run_cli(self.chapter_command())
        blocks = list(iter_block_items(Document(self.output)))
        tables = [(i, block) for i, block in enumerate(blocks) if isinstance(block, Table) and block.cell(0, 0).text == '编号']
        self.assertEqual(len(tables), 3)
        values = []
        for index, table in tables:
            self.assertTrue(blocks[index - 1].text.startswith('表2-1'))
            self.assertTrue(blocks[index - 1].paragraph_format.keep_with_next)
            self.assertTrue(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'))
            values.extend(row.cells[0].text for row in table.rows[1:])
        self.assertEqual(values, [str(i) for i in range(1, 28)])
        self.assertEqual(blocks[tables[0][0] - 2].text, '')
        self.assertEqual(blocks[tables[-1][0] + 1].text, '注：表下注释')
        self.assertEqual(blocks[tables[-1][0] + 2].text, '')
        self.assertEqual(blocks[tables[1][0] - 1].text, '表2-1（续表） 长表')

    def test_markdown_handles_escaped_pipes_and_rejects_ragged_tables(self):
        self.write_chapter('# 第2章 管线\n\n编号 | 名称\n--- | :---:\n1 | A\\|B\n')
        blocks = parse_markdown(self.markdown)
        self.assertEqual(blocks[-1].rows, [['编号', '名称'], ['1', 'A|B']])
        self.write_chapter('# 第2章 管线\n\n编号 | 名称\n--- | ---\n1 | A | extra\n')
        with self.assertRaisesRegex(ValueError, 'columns'):
            parse_markdown(self.markdown)

    def test_chapter_structure_and_unsupported_content_fail_before_output(self):
        invalid = [
            '前置说明\n\n# 第2章 测试\n',
            '# 第2章 测试\n\n# 第3章 错误\n',
            '# 第2章 测试\n\n## 3.1 错误\n',
            '# 第2章 测试\n\n## 2.1 正确\n\n## 2.1 重复\n',
            '# 第2章 测试\n\n![图](x.png)\n',
            '# 第2章 测试\n\n```python\nprint(1)\n```\n',
        ]
        for draft in invalid:
            with self.subTest(draft=draft):
                self.write_chapter(draft)
                with self.assertRaises(ValueError):
                    parse_markdown(self.markdown)

    def test_only_the_donors_used_in_this_chapter_are_required(self):
        self.write_chapter('# 第2章 测试\n\n新正文。\n')
        self.assertEqual(required_donor_keys(parse_markdown(self.markdown)), ['body', 'heading1'])
        document = self.make_document()
        donors = find_heading_donors(document, required=['body', 'heading1'])
        self.assertEqual(donors['body'].text, '真实第一章正文，应一直保留。')

    def test_style_donor_preserves_east_asian_font_and_does_not_clone_section_breaks(self):
        document = Document()
        donor = document.add_paragraph('1  标题')
        set_east_asia_font(donor.runs[0], east_asia_font='黑体', ascii_font='Times New Roman', size_pt=14)
        donor._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
        inserted = insert_paragraph_after(donor, donor, '2  新标题')
        fonts = inserted.runs[0]._r.rPr.rFonts
        self.assertEqual(fonts.get(qn('w:eastAsia')), '黑体')
        self.assertEqual(fonts.get(qn('w:ascii')), 'Times New Roman')
        self.assertFalse(inserted._p.xpath('./w:pPr/w:sectPr'))
        self.assertTrue(donor._p.xpath('./w:pPr/w:sectPr'))

    def test_verified_template_trim_removes_tables_and_content_controls(self):
        document = self.make_document(all_donors=True)
        content_control = OxmlElement('w:sdt')
        content = OxmlElement('w:sdtContent')
        content.append(OxmlElement('w:p'))
        content_control.append(content)
        document.paragraphs[4]._p.addnext(content_control)
        document.save(self.source)
        template = self.root / 'official.docx'
        shutil.copyfile(self.source, template)
        strip_template_body_placeholders(self.source, template, self.output)
        output = Document(self.output)
        self.assertEqual(len(output.tables), 0)
        self.assertFalse(output.element.xpath('.//w:sdt'))
        self.assertIn('参考文献原条目', [p.text for p in output.paragraphs])
        self.assertNotIn('真实第一章正文，应一直保留。', [p.text for p in output.paragraphs])
        changed = Document(self.source)
        changed.paragraphs[2].text = '用户已经改过的真实正文'
        changed.save(self.source)
        with self.assertRaisesRegex(RuntimeError, '不完全一致'):
            strip_template_body_placeholders(self.source, template, self.root / 'unsafe.docx')

    def test_missing_boundary_duplicate_headings_and_complex_fields_fail_closed(self):
        self.write_chapter()
        for case in ('missing', 'duplicate', 'field', 'revision'):
            with self.subTest(case=case):
                document = self.make_document()
                if case == 'missing':
                    delete_range(document, document.paragraphs[8], None)
                elif case == 'duplicate':
                    document.add_paragraph('2  重复章', style='Heading 1')
                elif case == 'field':
                    run = document.paragraphs[5].add_run()
                    field = OxmlElement('w:fldChar')
                    field.set(qn('w:fldCharType'), 'begin')
                    run._r.append(field)
                else:
                    document.paragraphs[5]._p.append(OxmlElement('w:ins'))
                document.save(self.source)
                self.run_cli(self.chapter_command(), expected=1)
                self.assertFalse(self.output.exists())

    def test_heading_override_mismatch_is_rejected(self):
        self.make_document()
        self.write_chapter()
        self.run_cli(self.pipeline_command() + ['--chapter-heading', '2  另一个标题'], expected=1)
        self.assertFalse(self.output.exists())

    def test_noop_only_copies_input_without_any_metadata_or_template(self):
        self.make_document()
        result = self.run_cli(self.pipeline_command(skip_chapter=True))
        self.assertEqual(self.source.read_bytes(), self.output.read_bytes())
        report = json.loads(result.stdout)
        self.assertEqual(report['steps_run'], [])
        self.assertFalse(report['contents_refresh_required'])

    def test_failed_later_stage_keeps_prior_output_and_recovers_completed_chapter(self):
        self.make_document()
        self.write_chapter()
        self.output.write_bytes(b'prior output must survive')
        original = self.source.read_bytes()
        references = self.root / 'empty_refs.md'
        references.write_text('# Empty reference list\n', encoding='utf-8')
        result = self.run_cli(self.pipeline_command(references=references), expected=1)
        report = json.loads(result.stderr[result.stderr.index('{'):])
        recovery_dir = Path(report['recovery_dir'])
        try:
            self.assertEqual(report['step'], 'references')
            self.assertFalse(report['output_published'])
            self.assertTrue(Path(report['last_successful_docx']).is_file())
            self.assertIn('修订后的第二章内容。', self.text_blocks(Path(report['last_successful_docx'])))
            self.assertEqual(self.output.read_bytes(), b'prior output must survive')
            self.assertEqual(self.source.read_bytes(), original)
        finally:
            shutil.rmtree(recovery_dir)


if __name__ == '__main__':
    unittest.main()
