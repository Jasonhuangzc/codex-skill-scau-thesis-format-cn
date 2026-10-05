import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document

from test_frontmatter import template, metadata

SCRIPTS = Path(__file__).resolve().parents[1] / "scau-thesis-format-cn" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from word_template_utils import ensure_output_copy


class OutputCopyTests(unittest.TestCase):
    def test_frontmatter_cli_rejects_overwriting_source_before_mutation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            docx = root / "working.docx"
            meta = root / "metadata.json"
            template().save(docx)
            meta.write_text(json.dumps(metadata()), encoding="utf-8")
            before = docx.read_bytes()
            result = subprocess.run([sys.executable, str(SCRIPTS / "fill_scau_frontmatter.py"),
                "--template", str(docx), "--meta", str(meta), "--output", str(docx)], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(docx.read_bytes(), before)

    def test_chapter_cli_rejects_overwriting_source_before_mutation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            docx = root / "working.docx"
            chapter = root / "chapter.md"
            doc = Document()
            doc.add_paragraph("3 结果", style="Heading 1")
            doc.add_paragraph("旧内容必须保留")
            doc.add_paragraph("参考文献", style="Heading 1")
            doc.save(docx)
            chapter.write_text("# 第3章 结果\n\n新内容", encoding="utf-8")
            before = docx.read_bytes()
            result = subprocess.run([sys.executable, str(SCRIPTS / "insert_markdown_chapter.py"),
                "--docx", str(docx), "--chapter-file", str(chapter), "--output", str(docx)], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(docx.read_bytes(), before)

    def test_hardlink_alias_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "source.docx"
            output = Path(td) / "alias.docx"
            source.write_bytes(b"source")
            os.link(source, output)
            with self.assertRaises(ValueError):
                ensure_output_copy(output, source)
            self.assertEqual(source.read_bytes(), b"source")


if __name__ == "__main__":
    unittest.main()
