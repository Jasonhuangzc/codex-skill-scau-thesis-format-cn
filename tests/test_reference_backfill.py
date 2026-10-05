import sys
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scau-thesis-format-cn" / "scripts"))
import insert_reference_batch as refs


class ReferenceBackfillTests(unittest.TestCase):
    def test_unstyled_toc_cannot_be_bibliography_replacement_anchor(self):
        doc = Document()
        field = doc.add_paragraph()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        field.add_run()._r.append(begin)
        instr = OxmlElement("w:instrText")
        instr.text = ' TOC \\o "1-3" '
        field.add_run()._r.append(instr)
        doc.add_paragraph("参考文献")
        end_para = doc.add_paragraph()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_para.add_run()._r.append(end)
        doc.add_paragraph("1 绪论", style="Heading 1")
        doc.add_paragraph("正文必须保留")
        heading = doc.add_paragraph("参 考 文 献", style="Heading 1")
        doc.add_paragraph("Brown A. Old reference[J]. 2024.")
        doc.add_paragraph("致 谢", style="Heading 1")
        self.assertIs(refs.find_reference_heading(doc)._p, heading._p)
        _, _, anchor = refs.find_reference_range(doc, heading)
        donor = refs.find_reference_donor(doc, heading, anchor)
        refs.validate_reference_range(doc, heading, anchor, reformat_only=False)
        refs.clear_existing_entries(doc, heading, anchor)
        refs.insert_entries(doc, heading, anchor, donor, ["Brown A. New reference[J]. 2026."])
        self.assertIn("正文必须保留", [p.text for p in doc.paragraphs])
        self.assertIn("参考文献", [p.text for p in doc.paragraphs])
        self.assertNotIn("Brown A. Old reference[J]. 2024.", [p.text for p in doc.paragraphs])

    def test_hyperlink_reformat_does_not_duplicate_visible_reference(self):
        doc = Document()
        heading = doc.add_paragraph("参考文献")
        entry = doc.add_paragraph("Author. ")
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Source title"
        run.append(text)
        link.append(run)
        entry._p.append(link)
        with self.assertRaisesRegex(RuntimeError, "complex objects"):
            refs.validate_reference_range(doc, heading, None, reformat_only=True)


if __name__ == "__main__":
    unittest.main()
