"""Regression tests use a synthetic template, never private thesis material."""
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scau-thesis-format-cn" / "scripts"))
import fill_scau_frontmatter as fm


def template():
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("本科毕业论文(或设计)")
    doc.add_paragraph("")
    doc.add_paragraph("论文（或设计）题目")
    for label, count in [("学    院:", 7), ("专    业:", 6), ("姓    名:", 6),
                         ("学    号:", 6), ("指导教师:", 10), ("提交日期：", 10)]:
        p = doc.add_paragraph()
        p.add_run(label)
        for _ in range(count - 1):
            p.add_run(" ")
    doc.add_paragraph("")
    section = doc.add_paragraph("")
    section._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    doc.add_paragraph("")
    doc.add_paragraph("华南农业大学本科毕业论文（设计）原创性声明")
    doc.add_paragraph("声明文字必须保留")
    doc.add_paragraph("使用授权声明")
    doc.add_paragraph("授权文字必须保留")
    doc.add_paragraph("摘        要")
    donor = doc.add_paragraph("中文摘要样例")
    marker = OxmlElement("w:bookmarkStart")
    marker.set(qn("w:id"), "9")
    marker.set(qn("w:name"), "abstract_anchor")
    donor._p.insert(0, marker)
    doc.add_paragraph("")
    doc.add_paragraph("关键词：关键词；关键词；关键词")
    doc.add_paragraph("English Title")
    doc.add_paragraph("Song Nianxiu")
    doc.add_paragraph("（College, University, Guangzhou 510642, China）")
    doc.add_paragraph("Abstract: Sample abstract.")
    doc.add_paragraph("")
    doc.add_paragraph("Key words: Sample; Sample; Sample")
    doc.add_paragraph("目        录")
    doc.add_paragraph("1 绪论", style="Heading 1")
    doc.add_paragraph("已有正文必须保留")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "已有数据必须保留"
    return doc


def metadata():
    return dict(thesis_title_zh="测试标题", college="测试学院", major="测试专业",
                student_name_zh="测试学生", student_id=20260001, advisor_name_zh="测试导师",
                advisor_title="教授", submission_date="2026年10月5日",
                thesis_title_en="Test Title", english_name="Test Student",
                college_en="Test College", university_en="Test University",
                city_en="Guangzhou", postal_code="510642", abstract_zh="第一段内容。\n\n第二段内容。\n第三段内容。",
                abstract_en="First paragraph.\n\nSecond paragraph.\nThird paragraph.",
                keywords_zh=["模板", "内容", "检查"], keywords_en=["Template", "Content", "Audit"])


def fill(doc, meta):
    fm.validate_template(doc)
    fm.replace_cover(doc, meta, "本科毕业论文")
    fm.replace_abstract_frontmatter(doc, meta)
    fm.normalize_cover_gap(doc)
    fm.ensure_frontmatter_page_breaks(doc)


class FrontmatterTests(unittest.TestCase):
    def test_repeat_fill_after_gap_cleanup_and_paragraph_growth(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "working.docx"
            doc = template()
            fill(doc, metadata())
            doc.save(path)
            doc = Document(path)
            count = len(doc.paragraphs)
            run_counts = [len(p.runs) for p in doc.paragraphs]
            expected = [p.text for p in doc.paragraphs]
            fill(doc, metadata())
            doc.save(path)
            reloaded = Document(path)
            self.assertEqual(len(reloaded.paragraphs), count)
            self.assertEqual([len(p.runs) for p in reloaded.paragraphs], run_counts)
            self.assertEqual([p.text for p in reloaded.paragraphs], expected)
            self.assertEqual(reloaded.tables[0].cell(0, 0).text, "已有数据必须保留")
            self.assertIn("已有正文必须保留", expected)
            self.assertTrue(reloaded.element.body.xpath(".//w:bookmarkStart[@w:name='abstract_anchor']"))
            self.assertEqual(len(reloaded.element.body.xpath(".//w:sectPr")), 2)
            fields = fm.resolve_frontmatter(reloaded)
            self.assertTrue(fields["thesis_title_en"].paragraph_format.page_break_before)
            self.assertEqual(fields["keywords_zh"].text, "关键词：模板；内容；检查")
            self.assertEqual(fields["keywords_en"].text, "Key words: Template; Content; Audit")
            self.assertTrue(fields["abstract_en"].runs[0].bold)
            self.assertFalse(fields["abstract_en"].runs[-1].bold)
            continuation = fields["abstract_en_paragraphs"][1]
            self.assertEqual(continuation._p.xpath("./w:pPr/w:ind/@w:firstLineChars"), ["200"])

    def test_omitted_optional_content_is_preserved(self):
        doc = template()
        fill(doc, metadata())
        fields = fm.resolve_frontmatter(doc)
        before = [p.text for p in fields["abstract_en_paragraphs"]]
        meta = metadata()
        for key in ("abstract_zh", "abstract_en", "keywords_zh", "keywords_en"):
            meta.pop(key)
        fill(doc, meta)
        self.assertEqual([p.text for p in fm.resolve_frontmatter(doc)["abstract_en_paragraphs"]], before)

    def test_ambiguous_anchor_rejected_before_writing(self):
        doc = template()
        doc.add_paragraph("关键词：重复锚点")
        before = doc.element.xml
        with self.assertRaises(RuntimeError):
            fm.validate_template(doc)
        self.assertEqual(doc.element.xml, before)

    def test_invalid_calendar_date_rejected(self):
        with self.assertRaises(ValueError):
            fm.split_submission_date("2026-02-30")

    def test_run_mapping_changes_are_rejected(self):
        doc = template()
        college = fm.resolve_frontmatter(doc)["college"]
        college.text = "学院：改版模板"
        with self.assertRaises(RuntimeError):
            fm.validate_template(doc)

    def test_complex_abstract_is_not_silently_erased(self):
        doc = template()
        abstract = fm.resolve_frontmatter(doc)["abstract_zh_paragraphs"][0]
        run = abstract.add_run()
        run._r.append(OxmlElement("w:drawing"))
        before = doc.element.xml
        with self.assertRaisesRegex(RuntimeError, "complex objects"):
            fm.validate_template(doc)
        self.assertEqual(doc.element.xml, before)

    def test_working_copy_comment_references_survive_text_updates(self):
        doc = template()
        donor = fm.resolve_frontmatter(doc)["abstract_zh_paragraphs"][0]
        ref = OxmlElement("w:commentReference")
        ref.set(qn("w:id"), "9")
        donor.add_run()._r.append(ref)
        fill(doc, metadata())
        fill(doc, metadata())
        self.assertEqual(len(doc.element.body.xpath(".//w:commentReference[@w:id='9']")), 1)

    def test_split_cover_field_fails_instead_of_retaining_old_suffix(self):
        doc = template()
        fill(doc, metadata())
        college = fm.resolve_frontmatter(doc)["college"]
        college.runs[6].text = "测试"
        college.add_run("学院")
        before = doc.element.xml
        with self.assertRaisesRegex(RuntimeError, "split"):
            fm.validate_template(doc)
        self.assertEqual(doc.element.xml, before)


if __name__ == "__main__":
    unittest.main()
